from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[3]
DOCX_PATH = ROOT / "运营管理" / "产品交付与宣传资料" / "学习悬浮助手操作手册.docx"


def insert_before(anchor: Paragraph, text: str, style: str) -> Paragraph:
    paragraph_xml = OxmlElement("w:p")
    anchor._p.addprevious(paragraph_xml)
    paragraph = Paragraph(paragraph_xml, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


document = Document(DOCX_PATH)

heading_text = "面试过程中临时答题或 Coding"
if not any(paragraph.text == heading_text for paragraph in document.paragraphs):
    anchor = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "08 / SETTINGS"
    )
    entries = [
        (heading_text, "Heading 2"),
        (
            "面试过程中临时遇到答题或手撕代码时，可以让面试助手继续听写，同时临时开启笔试助手截图搜题。切换前必须先把笔试助手调整为“悬浮框文字模式”。",
            "Normal",
        ),
        ("1｜保持面试助手运行", "Normal"),
        (
            "Alt + R 启动的面试听写和面试悬浮框无需关闭。",
            "Normal",
        ),
        ("2｜切换笔试助手模式", "Normal"),
        (
            "回到工作台进入笔试助手，选择“悬浮框文字模式”，不要使用语音播报模式。",
            "Normal",
        ),
        ("3｜启动笔试悬浮框", "Normal"),
        ("按 Alt + B 显示笔试助手悬浮框。", "Normal"),
        ("4｜截图并搜题", "Normal"),
        (
            "按 Alt + Q 截图，再按 Alt + E 搜题。Coding 题同样通过截图提交，答案会显示在笔试助手悬浮框中。",
            "Normal",
        ),
        ("5｜完成后关闭笔试悬浮框", "Normal"),
        (
            "再次按 Alt + B 隐藏笔试助手悬浮框，面试听写会继续运行；面试结束时按 Alt + R。",
            "Normal",
        ),
        (
            "两个助手可以同时运行。后启动或最近显示的悬浮框会位于上层；位置、宽度和高度调节会作用于当前最上层的悬浮框。",
            "Normal",
        ),
        ("", "Normal"),
    ]
    for text, style in entries:
        inserted = insert_before(anchor, text, style)
        if style == "Normal":
            inserted.paragraph_format.keep_with_next = False

document.save(DOCX_PATH)

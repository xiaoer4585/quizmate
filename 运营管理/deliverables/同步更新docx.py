# -*- coding: utf-8 -*-
"""
同步更新 2 份小红书 .docx（秋招第 6 篇 + 第 7 篇）
- 仅含标题 + 正文（文档清洁规则）
- 反问模式 + 避免敏感词
- 只附录官网一次
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

CN = "微软雅黑"
DARK = RGBColor(0x1A, 0x1A, 0x1A)
INK = RGBColor(0x2D, 0x2D, 0x2D)


def set_cn(run, name=CN, size=None, bold=None, color=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_paragraph_with_breaks(doc, text, size=12):
    blocks = [b for b in text.split("\n\n") if b.strip()]
    for b in blocks:
        lines = b.split("\n")
        p = doc.add_paragraph()
        for i, ln in enumerate(lines):
            if i > 0:
                p.add_run().add_break()
            run = p.add_run(ln)
            set_cn(run, size=size, color=INK)
        p.paragraph_format.space_after = Pt(6)


def make_doc(title, body, out_path):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = CN
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN)

    t = doc.add_heading(title, level=0)
    for run in t.runs:
        set_cn(run, size=22, bold=True, color=DARK)

    add_paragraph_with_breaks(doc, body, size=12)
    doc.save(out_path)
    print("saved", out_path)


# 第 6 篇
title6 = "秋招刷题神器不切屏了📚"
body6 = """秋招er 是不是都这样：白天投简历、晚上刷题、地铁上做公基、午饭时背行测公式——

时间已经被切成了碎片，可你还在用"截屏 → 切 App → 搜 → 看答案 → 切回 → 选"这套老流程刷题？
可以吗？×，不行，当然会被发现——而且思路断了七八次，1 题就吃掉 30 秒，一晚上根本没刷几套😭

用浏览器插件在手机上可以吗？×，不行，手机哪有浏览器插件。
用夸克悬浮球可以吗？×，不行，那不是干这用的。

那秋招er 到底怎么办？🤔

这里有解决方案——

按一下，AI 自动读题 + 分析 + 给答案，思路不打断。
新人可以送免费注册 + 送 50 积分，先体验再决定。
具体可以看官网：https://www.quizmate.vip/"""

# 第 7 篇
title7 = "秋招行测提速靠这招💡"
body7 = """行测刷到怀疑人生的姐妹兄弟看过来🙋‍♀️🙋‍♂️

题量大、时间紧、题型杂——常识/言语/数量/判断/资料分析，五大模块轮番暴击。
做完题对答案，效率比裸考还低，是常态吧？

靠"题海战术"，可以吗？×，不行，60 秒/题根本来不及。
靠"猜题押题"，可以吗？×，不行，秋招每年都换。

那到底靠什么？🤔

这里有解决方案——

不切屏、不截屏，半透明面板贴在题目旁边，AI 自动读题+分析+给答案。
你的资料能整本传上去，AI 答题时优先用你的资料。
新人可以送免费注册 + 送 50 积分，先体验再决定。
具体可以看官网：https://www.quizmate.vip/"""


def main():
    out6 = Path(r"E:\ai项目\考试插件\运营管理\deliverables\2026-07-13-秋招推广\小红书-第6篇-秋招刷题神器.docx")
    out7 = Path(r"E:\ai项目\考试插件\运营管理\deliverables\2026-07-13-秋招推广\小红书-第7篇-秋招行测提速.docx")
    make_doc(title6, body6, out6)
    make_doc(title7, body7, out7)


if __name__ == "__main__":
    main()

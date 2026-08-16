# -*- coding: utf-8 -*-
"""
头条号 + 百家号 · Word 文档生成器

依据：
- 头条号：只支持 Word 导入（强制）
- 百家号：只支持 Word 导入 + 含表格强制 Word

合规策略：
- 头条号（🟠 严）：去品牌化（不出现产品名）+ 无外链 + 短平快
- 百家号（🟢 松）：可保留产品名 + 1 个官网链接 + 结构化 + 必含表格
"""
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from topics_data import TOPICS, WEBSITE_URL

# python-docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 配置
BASE_DIR = r"E:\ai项目\考试插件\运营管理\deliverables\2026-07-25-秋招多平台20篇-md"
BRAND_RED = RGBColor(0xFF, 0x24, 0x42)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
GRAY_TEXT = RGBColor(0x88, 0x88, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CN_FONT = "微软雅黑"


# ========== 工具函数 ==========
def set_cn(run, size=11, bold=False, color=None):
    run.font.name = CN_FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), CN_FONT)
    rFonts.set(qn('w:ascii'), CN_FONT)
    rFonts.set(qn('w:hAnsi'), CN_FONT)


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def make_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = CN_FONT
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), CN_FONT)
    rfonts.set(qn('w:ascii'), CN_FONT)
    rfonts.set(qn('w:hAnsi'), CN_FONT)
    return doc


def add_title(doc, text, size=20, color=BRAND_RED, center=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_cn(run, size=size, bold=True, color=color)
    p.paragraph_format.space_after = Pt(12)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run("▎" + text)
    set_cn(run, size=14, bold=True, color=BRAND_RED)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_para(doc, text, size=11, bold=False, color=None, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.6
    run = p.add_run(text)
    set_cn(run, size=size, bold=bold, color=color or DARK_TEXT)
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(text)
    set_cn(run, size=size, color=DARK_TEXT)
    return p


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade_cell(cell, "FF2442")
        run = cell.paragraphs[0].add_run(h)
        set_cn(run, size=11, bold=True, color=WHITE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            run = cells[i].paragraphs[0].add_run(str(val))
            set_cn(run, size=10)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    return t


def add_cta(doc, text=None):
    if text is None:
        text = f"📚 完整功能与下载：{WEBSITE_URL}"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run(text)
    set_cn(run, size=11, color=BRAND_RED, bold=True)
    return p


def save_doc(doc, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)


# ========== 头条号（短平快 + 蹭热点，无品牌） ==========
def render_toutiao(topic):
    """头条号：短平快 + 蹭热点 + 去品牌化"""
    doc = make_doc()

    title = topic['title_other']
    # 标题去掉耸人听闻
    title = title.replace("错过再等一年", "完整梳理")
    add_title(doc, title, size=20)
    add_para(doc, f"{topic['category']} · 阅读约 4 分钟",
             color=GRAY_TEXT, size=10, indent=False)

    # 开篇点题
    add_para(doc, "【开篇点题】", bold=True)
    add_para(doc, topic["core"])

    # 快速清单
    add_h2(doc, "快速清单")
    for f in topic["facts"]:
        add_para(doc, f"✅ {f}", indent=False)

    # 核心方法
    add_h2(doc, "核心方法")
    tip = topic["tip"]
    # 头条号：去产品名 + 去"工具"硬广
    for kw in ["求职雷达", "答题悬浮助手", "quizmate.vip", "watch.quizmate.vip"]:
        tip = tip.replace(kw, "相应工具")
    tip = tip.replace("推荐用网页答题工具", "建议结合合适方式")
    tip = tip.replace("推荐用相应工具", "建议结合合适方式")
    tip = tip.replace("用网页答题工具", "通过合适的方式")
    tip = tip.replace("用相应工具", "通过合适的方式")
    tip = tip.replace("用招聘流程监控工具", "通过合适的方式")
    add_para(doc, tip, bold=True, color=BRAND_RED)

    add_para(doc, "把这三件事做扎实：明确目标 → 建立信息源 → 持续跟进。", indent=False)

    # 为什么重要
    add_h2(doc, "为什么这件事秋招时一定要做")
    add_para(doc, "1. 把握时间窗口：秋招只有 2 个月主战场。", indent=False)
    add_para(doc, "2. 避免漏掉机会：分散招聘渠道容易错过公告。", indent=False)
    add_para(doc, "3. 提升通过率：节奏感比刷题量更关键。", indent=False)

    # 经验提醒
    add_h2(doc, "经验提醒")
    add_para(doc, "海投简历不看 JD、通过率暴跌，是秋招最常见的坑；等到 9 月才开始准备，提前批已经结束；只刷题不总结、错题反复错。避开这三点，通过率显著提升。")

    # 写在最后
    add_h2(doc, "写在最后")
    add_para(doc, "秋招信息分散是关键痛点，建议建立一张信息源清单 + 投递进度跟踪表，把节奏感视觉化。如果你有其他问题，欢迎评论区留言讨论。",
             indent=False)

    save_doc(doc, os.path.join(BASE_DIR, "05_头条号", f"{topic['id']:02d}_{topic['topic']}.docx"))
    return True


# ========== 百家号（结构化 + 表格 + 可保留品牌） ==========
def render_baidu(topic):
    """百家号：结构化 + 必含表格 + 1 个官网链接（可保留品牌名）"""
    doc = make_doc()

    title = topic['title_other']
    title = title.replace("错过再等一年", "完整梳理")
    add_title(doc, title, size=20)
    add_para(doc, f"主题：{topic['topic']} · 分类：{topic['category']}",
             color=GRAY_TEXT, size=10, indent=False)

    # 导语
    add_para(doc, "导语：", bold=True)
    add_para(doc, topic["core"])

    # 本文要点
    add_h2(doc, "一、本文要点")
    for f in topic["facts"]:
        add_bullet(doc, f)

    # 详细对比表（必含）
    add_h2(doc, "二、详细对比表")
    if topic["category"] == "时间点":
        add_table(doc,
                  ["批次", "时间窗口", "适用对象", "建议动作"],
                  [
                      ["提前批", "5-8 月", "顶尖学生/技术岗", "大胆投递练手"],
                      ["正式批", "8-10 月", "应届生主力", "集中精力主攻"],
                      ["补录批", "11-12 月", "未招满岗位", "跟进进度 + 补投"],
                      ["春招", "次年 3-5 月", "秋招失利者", "复盘后二次尝试"],
                  ])
    elif topic["category"] == "流程监控":
        add_table(doc,
                  ["阶段", "动作", "频率"],
                  [
                      ["信息获取", "浏览新岗位", "每日"],
                      ["简历投递", "定制化投递", "每周 5-10 个"],
                      ["笔试准备", "刷题训练", "每天 1 套"],
                      ["面试准备", "项目复盘", "面试前 3 天"],
                  ])
    elif topic["category"] == "行测题":
        add_table(doc,
                  ["模块", "题量占比", "难度", "训练建议"],
                  [
                      ["言语理解", "约 25%", "中等", "每天 30 题"],
                      ["数量关系", "约 20%", "较难", "按题型集中刷"],
                      ["判断推理", "约 25%", "中等", "分题型速解"],
                      ["资料分析", "约 20%", "易提分", "每天 5 篇"],
                      ["常识判断", "约 10%", "积累型", "碎片时间记忆"],
                  ])
    elif topic["category"] == "题库":
        add_table(doc,
                  ["模块", "题量建议", "训练周期", "检验标准"],
                  [
                      ["基础题型", "200 题", "1 周", "正确率 80%"],
                      ["进阶题型", "100 题", "1 周", "正确率 70%"],
                      ["综合训练", "50 套卷", "2 周", "平均 60 分"],
                      ["错题归档", "全部错题", "持续", "复盘周期 7 天"],
                  ])
    else:  # 笔试/秋招
        add_table(doc,
                  ["环节", "占总成绩", "应对策略"],
                  [
                      ["通用行测", "30%", "刷题 + 限速"],
                      ["专业题", "50%", "按岗位准备"],
                      ["综合开放题", "20%", "结构化答题"],
                  ])

    # 实操步骤
    add_h2(doc, "三、实操步骤")
    tip = topic["tip"]
    # 百家号：可保留品牌名（账号是教育类），但只 1 个官网链接
    add_para(doc, f"核心建议：{tip}", bold=True, color=BRAND_RED)
    add_para(doc, "Step 1：明确目标与方向。", indent=False)
    add_para(doc, "Step 2：建立信息源清单与跟踪表。", indent=False)
    add_para(doc, "Step 3：模块化刷题与定期复盘。", indent=False)
    add_para(doc, "Step 4：持续跟进投递进度。", indent=False)
    add_para(doc, "Step 5：每周日做整体复盘。", indent=False)

    # 常见误区
    add_h2(doc, "四、常见误区")
    add_para(doc, "❌ 海投简历不看 JD：通过率暴跌，HR 容易把你拉黑", indent=False)
    add_para(doc, "❌ 等到 9 月才开始准备：提前批已结束，错过一半机会", indent=False)
    add_para(doc, "❌ 只刷题不总结：题海战术浪费时间，错题反复错", indent=False)
    add_para(doc, "❌ 投完就忘：没有进度跟踪表等于盲投", indent=False)

    # 推荐工具（百家号可保留品牌名）
    add_h2(doc, "五、推荐工具")
    add_para(doc, "求职雷达（监控招聘流程）+ 答题悬浮助手（网页答题 + 个人知识库）。配合使用覆盖秋招全流程。", indent=False)

    add_cta(doc)

    save_doc(doc, os.path.join(BASE_DIR, "07_百家号", f"{topic['id']:02d}_{topic['topic']}.docx"))
    return True


# ========== 主调度 ==========
def main():
    print("="*60)
    print("头条号 + 百家号 · Word 文档生成")
    print("="*60)

    print("\n[05_头条号] 生成 20 篇 .docx（去品牌化 + 无外链）")
    for topic in TOPICS:
        try:
            render_toutiao(topic)
            print(f"  ✓ {topic['id']:02d}. {topic['topic']}")
        except Exception as e:
            print(f"  ✗ {topic['id']:02d}. {topic['topic']} - {e}")

    print("\n[07_百家号] 生成 20 篇 .docx（含表格 + 1 个官网链接）")
    for topic in TOPICS:
        try:
            render_baidu(topic)
            print(f"  ✓ {topic['id']:02d}. {topic['topic']}")
        except Exception as e:
            print(f"  ✗ {topic['id']:02d}. {topic['topic']} - {e}")

    print("\n✓ 全部 40 篇 .docx 已生成")


if __name__ == "__main__":
    main()